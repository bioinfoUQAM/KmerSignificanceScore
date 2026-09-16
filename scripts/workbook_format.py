"""Shared presentation for the supplementary workbooks."""

from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

MAX_WIDTH = 52


def format_sheets(writer, max_width: int = MAX_WIDTH) -> None:
    """Make every sheet of a workbook readable without resizing anything by hand.

    Columns default to a width that truncates headers and runs numbers into each other, which
    is how a supplementary table reaches a reviewer unreadable. Each column is widened to its
    longest cell, the header row is bolded and frozen, and the width is capped so that one long
    text column cannot push the rest off the screen.

    Args:
        writer: an open pandas ExcelWriter using the openpyxl engine.
        max_width: widest a column may become, in characters.
    """
    for sheet in writer.book.worksheets:
        for column in sheet.columns:
            longest = max((len(str(cell.value)) for cell in column if cell.value is not None),
                          default=0)
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = min(max(longest + 2, 10), max_width)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        # Data cells wrap too. Wrapping the header alone was not enough: the width cap above
        # means a long cell is wider than its column, so without this it is simply cut off at
        # the border. S10 lost the tail of its HIV-1 class description and the justifications
        # in S12's admissibility sheet ran off the edge, both of them text a reader needs.
        # Row heights are left unset so the reader's spreadsheet fits them to the wrapped text.
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
